"""
Service for exporting and importing application data in various formats.
"""
import os
import json
import logging
from datetime import datetime
from typing import Any, Dict, List, Optional, Union, BinaryIO
from pathlib import Path
import io

from fastapi import HTTPException, status, UploadFile, File, Depends
from sqlalchemy.ext.asyncio import AsyncSession
from sqlalchemy.future import select
import pandas as pd
from docx import Document
from docx.shared import Inches
from reportlab.lib.pagesizes import letter
from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, Table, TableStyle, Image as RLImage
from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
from reportlab.lib import colors
from reportlab.lib.units import inch

from app.core.config import settings
from app.database import get_db
from app.models_user import User
from app.models import Research
from app.models_story import NameStory, OralHistory

logger = logging.getLogger(__name__)

class ExportImportService:
    """Service for handling data export and import operations."""
    
    def __init__(self, db: AsyncSession):
        """Initialize the export/import service.
        
        Args:
            db: Database session
        """
        self.db = db
        self.export_dir = os.path.join(settings.EXPORT_DIR, "exports")
        os.makedirs(self.export_dir, exist_ok=True)
    
    async def export_research(
        self, 
        research_id: int, 
        user_id: int,
        format: str = "json"
    ) -> str:
        """Export research data in the specified format.
        
        Args:
            research_id: ID of the research to export
            user_id: ID of the user requesting the export
            format: Export format (json, csv, pdf, docx)
            
        Returns:
            Path to the exported file
            
        Raises:
            HTTPException: If the research is not found or the user doesn't have permission
        """
        # Get the research with related data
        result = await self.db.execute(
            select(Research)
            .where(Research.id == research_id)
            .options(
                selectinload(Research.stories),
                selectinload(Research.oral_histories)
            )
        )
        research = result.scalars().first()
        
        if not research:
            raise HTTPException(
                status_code=status.HTTP_404_NOT_FOUND,
                detail="Research not found"
            )
        
        if research.user_id != user_id:
            raise HTTPException(
                status_code=status.HTTP_403_FORBIDDEN,
                detail="Not authorized to export this research"
            )
        
        # Prepare data for export
        export_data = {
            "research": {
                "id": research.id,
                "name": research.name,
                "status": research.status,
                "created_at": research.created_at.isoformat(),
                "updated_at": research.updated_at.isoformat() if research.updated_at else None,
            },
            "stories": [
                {
                    "id": story.id,
                    "title": story.title,
                    "content": story.content,
                    "source": story.source,
                    "created_at": story.created_at.isoformat(),
                    "updated_at": story.updated_at.isoformat() if story.updated_at else None,
                }
                for story in research.stories
            ],
            "oral_histories": [
                {
                    "id": oh.id,
                    "title": oh.title,
                    "description": oh.description,
                    "audio_path": oh.audio_path,
                    "transcript": oh.transcript,
                    "created_at": oh.created_at.isoformat(),
                    "updated_at": oh.updated_at.isoformat() if oh.updated_at else None,
                }
                for oh in research.oral_histories
            ]
        }
        
        # Export in the requested format
        filename = f"research_export_{research_id}_{datetime.now().strftime('%Y%m%d_%H%M%S')}"
        
        if format == "json":
            return self._export_to_json(export_data, filename)
        elif format == "csv":
            return self._export_to_csv(export_data, filename)
        elif format == "pdf":
            return self._export_to_pdf(export_data, filename)
        elif format == "docx":
            return self._export_to_docx(export_data, filename)
        else:
            raise HTTPException(
                status_code=status.HTTP_400_BAD_REQUEST,
                detail=f"Unsupported export format: {format}"
            )
    
    def _export_to_json(self, data: Dict, filename: str) -> str:
        """Export data to JSON file.
        
        Args:
            data: Data to export
            filename: Base filename (without extension)
            
        Returns:
            Path to the exported file
        """
        filepath = os.path.join(self.export_dir, f"{filename}.json")
        with open(filepath, 'w', encoding='utf-8') as f:
            json.dump(data, f, indent=2, ensure_ascii=False)
        return filepath
    
    def _export_to_csv(self, data: Dict, filename: str) -> str:
        """Export data to CSV files (one per entity type).
        
        Args:
            data: Data to export
            filename: Base filename (without extension)
            
        Returns:
            Path to the zip file containing CSV files
        """
        import zipfile
        import tempfile
        
        # Create a temporary directory for CSV files
        with tempfile.TemporaryDirectory() as temp_dir:
            # Export each entity type to a separate CSV
            for entity_type, items in data.items():
                if not items:
                    continue
                    
                df = pd.DataFrame(items)
                csv_path = os.path.join(temp_dir, f"{entity_type}.csv")
                df.to_csv(csv_path, index=False, encoding='utf-8')
            
            # Create a zip file
            zip_path = os.path.join(self.export_dir, f"{filename}.zip")
            with zipfile.ZipFile(zip_path, 'w') as zipf:
                for root, _, files in os.walk(temp_dir):
                    for file in files:
                        file_path = os.path.join(root, file)
                        arcname = os.path.relpath(file_path, temp_dir)
                        zipf.write(file_path, arcname)
        
        return zip_path
    
    def _export_to_pdf(self, data: Dict, filename: str) -> str:
        """Export data to a formatted PDF document.
        
        Args:
            data: Data to export
            filename: Base filename (without extension)
            
        Returns:
            Path to the exported PDF file
        """
        from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
        from reportlab.lib import colors
        from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, Table, TableStyle
        
        filepath = os.path.join(self.export_dir, f"{filename}.pdf")
        doc = SimpleDocTemplate(filepath, pagesize=letter)
        
        # Define styles
        styles = getSampleStyleSheet()
        title_style = ParagraphStyle(
            'Title',
            parent=styles['Heading1'],
            fontSize=16,
            spaceAfter=12,
            textColor=colors.HexColor('#2c3e50')
        )
        
        heading_style = ParagraphStyle(
            'Heading2',
            parent=styles['Heading2'],
            fontSize=14,
            spaceAfter=6,
            textColor=colors.HexColor('#3498db')
        )
        
        normal_style = styles['Normal']
        
        # Build the document
        story = []
        
        # Add title
        story.append(Paragraph(f"Research Export: {data['research']['name']}", title_style))
        story.append(Spacer(1, 12))
        
        # Add research details
        story.append(Paragraph("Research Details", heading_style))
        research_table_data = [
            ["Name", data['research']['name']],
            ["Status", data['research']['status']],
            ["Created", data['research']['created_at']],
        ]
        
        if data['research']['updated_at']:
            research_table_data.append(["Last Updated", data['research']['updated_at']])
        
        research_table = Table(research_table_data, colWidths=[100, 400])
        research_table.setStyle(TableStyle([
            ('VALIGN', (0, 0), (-1, -1), 'TOP'),
            ('FONTNAME', (0, 0), (-1, -1), 'Helvetica'),
            ('FONTSIZE', (0, 0), (-1, -1), 10),
            ('BOTTOMPADDING', (0, 0), (-1, -1), 6),
            ('TOPPADDING', (0, 0), (-1, -1), 6),
            ('GRID', (0, 0), (-1, -1), 0.5, colors.lightgrey),
        ]))
        
        story.append(research_table)
        story.append(Spacer(1, 24))
        
        # Add stories
        if data.get('stories'):
            story.append(Paragraph("Name Stories", heading_style))
            for i, s in enumerate(data['stories'], 1):
                story.append(Paragraph(f"{i}. {s['title']}", styles['Heading3']))
                story.append(Paragraph(s['content'], normal_style))
                if s.get('source'):
                    story.append(Paragraph(f"<i>Source: {s['source']}</i>", styles['Italic']))
                story.append(Spacer(1, 12))
        
        # Add oral histories
        if data.get('oral_histories'):
            story.append(Paragraph("Oral Histories", heading_style))
            for i, oh in enumerate(data['oral_histories'], 1):
                story.append(Paragraph(f"{i}. {oh['title']}", styles['Heading3']))
                if oh.get('description'):
                    story.append(Paragraph(oh['description'], normal_style))
                if oh.get('transcript'):
                    story.append(Paragraph("<b>Transcript:</b>", styles['Normal']))
                    story.append(Paragraph(oh['transcript'], normal_style))
                story.append(Spacer(1, 12))
        
        # Build the PDF
        doc.build(story)
        return filepath
    
    def _export_to_docx(self, data: Dict, filename: str) -> str:
        """Export data to a formatted Word document.
        
        Args:
            data: Data to export
            filename: Base filename (without extension)
            
        Returns:
            Path to the exported DOCX file
        """
        filepath = os.path.join(self.export_dir, f"{filename}.docx")
        doc = Document()
        
        # Add title
        doc.add_heading(f"Research Export: {data['research']['name']}", 0)
        
        # Add research details
        doc.add_heading('Research Details', level=1)
        research_table = doc.add_table(rows=1, cols=2)
        research_table.style = 'Table Grid'
        
        # Add header row
        hdr_cells = research_table.rows[0].cells
        hdr_cells[0].text = 'Field'
        hdr_cells[1].text = 'Value'
        
        # Add data rows
        rows = [
            ('Name', data['research']['name']),
            ('Status', data['research']['status']),
            ('Created', data['research']['created_at']),
        ]
        
        if data['research']['updated_at']:
            rows.append(('Last Updated', data['research']['updated_at']))
        
        for field, value in rows:
            row_cells = research_table.add_row().cells
            row_cells[0].text = field
            row_cells[1].text = str(value)
        
        # Add stories
        if data.get('stories'):
            doc.add_heading('Name Stories', level=1)
            for i, s in enumerate(data['stories'], 1):
                doc.add_heading(f"{i}. {s['title']}", level=2)
                doc.add_paragraph(s['content'])
                if s.get('source'):
                    p = doc.add_paragraph()
                    p.add_run('Source: ').italic = True
                    p.add_run(s['source'])
                doc.add_paragraph()
        
        # Add oral histories
        if data.get('oral_histories'):
            doc.add_heading('Oral Histories', level=1)
            for i, oh in enumerate(data['oral_histories'], 1):
                doc.add_heading(f"{i}. {oh['title']}", level=2)
                if oh.get('description'):
                    doc.add_paragraph(oh['description'])
                if oh.get('transcript'):
                    doc.add_heading('Transcript', level=3)
                    doc.add_paragraph(oh['transcript'])
                doc.add_paragraph()
        
        # Save the document
        doc.save(filepath)
        return filepath
    
    async def import_research(
        self, 
        file: UploadFile, 
        user_id: int,
        format: str = "json"
    ) -> Dict[str, Any]:
        """Import research data from a file.
        
        Args:
            file: Uploaded file containing the data
            user_id: ID of the user importing the data
            format: Import format (json, csv, pdf, docx)
            
        Returns:
            Dictionary with import status and imported research ID
            
        Raises:
            HTTPException: If the import fails
        """
        try:
            # Read file content
            content = await file.read()
            
            if format == "json":
                data = json.loads(content)
            elif format == "csv":
                # For CSV, we expect a ZIP file with multiple CSVs
                data = self._import_from_csv(content)
            elif format == "pdf":
                # For PDF, we can only extract text, not structure
                data = self._import_from_pdf(content)
            elif format == "docx":
                data = self._import_from_docx(content)
            else:
                raise HTTPException(
                    status_code=status.HTTP_400_BAD_REQUEST,
                    detail=f"Unsupported import format: {format}"
                )
            
            # Validate the data
            if not isinstance(data, dict) or 'research' not in data:
                raise HTTPException(
                    status_code=status.HTTP_400_BAD_REQUEST,
                    detail="Invalid data format: missing 'research' section"
                )
            
            # Create the research
            research_data = data['research']
            research = Research(
                name=research_data.get('name', 'Imported Research'),
                status=research_data.get('status', 'completed'),
                user_id=user_id
            )
            
            self.db.add(research)
            await self.db.commit()
            await self.db.refresh(research)
            
            # Add stories
            if 'stories' in data and isinstance(data['stories'], list):
                for story_data in data['stories']:
                    story = NameStory(
                        title=story_data.get('title', 'Untitled Story'),
                        content=story_data.get('content', ''),
                        source=story_data.get('source'),
                        research_id=research.id,
                        user_id=user_id
                    )
                    self.db.add(story)
            
            # Add oral histories
            if 'oral_histories' in data and isinstance(data['oral_histories'], list):
                for oh_data in data['oral_histories']:
                    oh = OralHistory(
                        title=oh_data.get('title', 'Untitled Oral History'),
                        description=oh_data.get('description'),
                        audio_path=oh_data.get('audio_path'),
                        transcript=oh_data.get('transcript'),
                        research_id=research.id,
                        user_id=user_id
                    )
                    self.db.add(oh)
            
            await self.db.commit()
            
            return {
                "status": "success",
                "message": "Research imported successfully",
                "research_id": research.id
            }
            
        except json.JSONDecodeError:
            raise HTTPException(
                status_code=status.HTTP_400_BAD_REQUEST,
                detail="Invalid JSON format"
            )
        except Exception as e:
            await self.db.rollback()
            logger.error(f"Error importing research: {str(e)}", exc_info=True)
            raise HTTPException(
                status_code=status.HTTP_500_INTERNAL_SERVER_ERROR,
                detail=f"Error importing research: {str(e)}"
            )
    
    def _import_from_csv(self, content: bytes) -> Dict[str, Any]:
        """Import data from CSV files in a ZIP archive.
        
        Args:
            content: ZIP file content containing CSV files
            
        Returns:
            Dictionary with the imported data
        """
        import zipfile
        import io
        import pandas as pd
        
        result = {}
        
        with zipfile.ZipFile(io.BytesIO(content)) as zip_ref:
            # Extract to a temporary directory
            with tempfile.TemporaryDirectory() as temp_dir:
                zip_ref.extractall(temp_dir)
                
                # Process each CSV file
                for root, _, files in os.walk(temp_dir):
                    for file in files:
                        if not file.lower().endswith('.csv'):
                            continue
                            
                        entity_type = os.path.splitext(file)[0].lower()
                        file_path = os.path.join(root, file)
                        
                        try:
                            df = pd.read_csv(file_path)
                            result[entity_type] = df.to_dict('records')
                        except Exception as e:
                            logger.warning(f"Error processing {file}: {str(e)}")
        
        return result
    
    def _import_from_pdf(self, content: bytes) -> Dict[str, Any]:
        """Extract text from a PDF file.
        
        Note: This is a very basic implementation that just extracts text.
        For a real application, you'd want to parse the structure.
        
        Args:
            content: PDF file content
            
        Returns:
            Dictionary with the extracted text
        """
        from PyPDF2 import PdfReader
        
        text = ""
        with io.BytesIO(content) as f:
            reader = PdfReader(f)
            for page in reader.pages:
                text += page.extract_text() + "\n\n"
        
        return {
            "research": {
                "name": "Imported from PDF",
                "status": "completed"
            },
            "stories": [
                {
                    "title": "Extracted Text",
                    "content": text,
                    "source": "PDF Import"
                }
            ]
        }
    
    def _import_from_docx(self, content: bytes) -> Dict[str, Any]:
        """Extract text from a Word document.
        
        Note: This is a very basic implementation that just extracts text.
        For a real application, you'd want to parse the structure.
        
        Args:
            content: DOCX file content
            
        Returns:
            Dictionary with the extracted text
        """
        from docx import Document as DocxDocument
        
        with io.BytesIO(content) as f:
            doc = DocxDocument(f)
            
            # Extract text from paragraphs
            text = "\n\n".join([p.text for p in doc.paragraphs if p.text.strip()])
            
            # Extract text from tables
            for table in doc.tables:
                for row in table.rows:
                    for cell in row.cells:
                        text += "\n" + "\n".join([p.text for p in cell.paragraphs if p.text.strip()])
        
        return {
            "research": {
                "name": "Imported from Word",
                "status": "completed"
            },
            "stories": [
                {
                    "title": "Extracted Text",
                    "content": text,
                    "source": "Word Import"
                }
            ]
        }

# Helper function to get the export/import service
def get_export_import_service(db: AsyncSession = Depends(get_db)):
    return ExportImportService(db)
